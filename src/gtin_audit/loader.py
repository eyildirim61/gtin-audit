"""Read CSV / XLSX exports and map their columns to the fields the audit needs."""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

from .checks import Record

FIELDS = ("ref", "name", "unit_barcode", "box_barcode", "units_per_box", "uom", "box_is_item")

# Header aliases (English, Turkish and common Odoo export names). Compared after normalisation.
ALIASES: Dict[str, Sequence[str]] = {
    "ref": (
        "ref", "default_code", "internal reference", "sku", "item code", "product code",
        "stok kodu", "urun kodu", "malzeme kodu", "kod",
    ),
    "name": (
        "name", "product", "product name", "description", "display name",
        "urun", "urun adi", "stok adi", "aciklama",
    ),
    "unit_barcode": (
        "unit_barcode", "unit barcode", "barcode", "ean", "ean13", "gtin", "upc",
        "barkod", "birim barkod", "birim barkodu", "urun barkodu",
    ),
    "box_barcode": (
        "box_barcode", "box barcode", "case barcode", "carton barcode", "packaging barcode",
        "packaging_ids/barcode", "packagings/barcode", "itf14", "gtin14",
        "koli barkod", "koli barkodu", "paket barkodu",
    ),
    "units_per_box": (
        "units_per_box", "units per box", "qty per box", "pack qty", "pack size", "case qty",
        "packaging_ids/qty", "packagings/contained quantity", "contained quantity",
        "koli ici adet", "koli adedi", "koli ici miktar", "paket ici adet",
    ),
    "uom": (
        "uom", "uom_id", "unit", "unit of measure", "units of measure", "sales unit",
        "birim", "olcu birimi", "satis birimi",
    ),
    "box_is_item": (
        "box_is_item", "box is item", "box = item", "one code", "single code",
        "same barcode for box", "box uses item barcode",
        "tek kod", "tek barkod", "koli = urun", "koli urun ayni barkod",
    ),
}

_TR_MAP = str.maketrans({"ı": "i", "ç": "c", "ş": "s", "ğ": "g", "ü": "u", "ö": "o"})


class LoaderError(Exception):
    """Raised when the input file cannot be read or mapped."""


def normalize_header(header: str) -> str:
    text = unicodedata.normalize("NFKD", (header or "").strip().casefold())
    text = "".join(c for c in text if not unicodedata.combining(c))  # drops the dot of İ
    text = text.translate(_TR_MAP)
    text = re.sub(r"[\s_\-.:]+", " ", text)
    return text.strip()


_PAREN_NOTE = re.compile(r"\s*\([^)]*\)\s*$")

_ALIAS_LOOKUP = {field: {normalize_header(a) for a in aliases} for field, aliases in ALIASES.items()}


def _cell_to_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    return str(value).strip()


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1254", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise LoaderError("Could not decode the CSV file.")  # pragma: no cover (latin-1 never fails)


def _read_csv(path: Path) -> Tuple[List[str], List[List[str]]]:
    text = _decode(path.read_bytes())
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), dialect)]
    if not rows:
        raise LoaderError(f"{path.name} is empty.")
    return rows[0], rows[1:]


def _read_xlsx(path: Path, sheet: Optional[str]) -> Tuple[List[str], List[List[str]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise LoaderError("Reading .xlsx needs openpyxl: pip install 'gtin-audit[xlsx]'") from exc

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet:
            if sheet not in wb.sheetnames:
                raise LoaderError(f"Sheet '{sheet}' not found. Sheets: {', '.join(wb.sheetnames)}")
            ws = wb[sheet]
        else:
            ws = wb.worksheets[0]
        rows = [[_cell_to_str(v) for v in row] for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    if not rows:
        raise LoaderError(f"{path.name} has no rows.")
    return rows[0], rows[1:]


def read_table(path, sheet: Optional[str] = None) -> Tuple[List[str], List[List[str]]]:
    path = Path(path)
    if not path.exists():
        raise LoaderError(f"File not found: {path}")
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        return _read_xlsx(path, sheet)
    if suffix in (".csv", ".tsv", ".txt"):
        return _read_csv(path)
    raise LoaderError(f"Unsupported file type '{suffix}'. Use .csv or .xlsx.")


def resolve_columns(headers: Sequence[str], overrides: Optional[Dict[str, str]] = None) -> Dict[str, int]:
    """Map each field to a column index, using explicit overrides first, then aliases."""
    overrides = {k: v for k, v in (overrides or {}).items() if v}
    normalized = [normalize_header(h) for h in headers]
    mapping: Dict[str, int] = {}

    for field, header in overrides.items():
        if field not in FIELDS:
            raise LoaderError(f"Unknown field '{field}'.")
        target = normalize_header(header)
        if target not in normalized:
            raise LoaderError(f"Column '{header}' not found. Columns: {', '.join(headers)}")
        mapping[field] = normalized.index(target)

    # Pass 1: exact header match. Pass 2: header with a trailing note in parentheses removed,
    # so "Barcode (EAN)" or "Same Barcode For Box (Yes/No)" match the plain aliases.
    stripped = [_PAREN_NOTE.sub("", name).strip() for name in normalized]
    for candidates in (normalized, stripped):
        for field in FIELDS:
            if field in mapping:
                continue
            for idx, name in enumerate(candidates):
                if name in _ALIAS_LOOKUP[field] and idx not in mapping.values():
                    mapping[field] = idx
                    break

    if "unit_barcode" not in mapping and "box_barcode" not in mapping:
        raise LoaderError(
            "Could not find a barcode column. Columns found: "
            + ", ".join(headers)
            + ". Use --unit-barcode-col / --box-barcode-col to name it."
        )
    return mapping


def load_records(
    path,
    sheet: Optional[str] = None,
    overrides: Optional[Dict[str, str]] = None,
    fill_down: bool = True,
) -> Tuple[List[Record], Dict[str, str]]:
    """Load a file into Records. Returns (records, {field: header used})."""
    headers, rows = read_table(path, sheet)
    mapping = resolve_columns(headers, overrides)

    def get(row: List[str], field: str) -> str:
        idx = mapping.get(field)
        if idx is None or idx >= len(row):
            return ""
        return row[idx].strip()

    records: List[Record] = []
    group = 0
    for i, row in enumerate(rows):
        if not any(c.strip() for c in row):
            continue
        values = {f: get(row, f) for f in FIELDS}
        is_continuation = (
            fill_down
            and records
            and not (values["ref"] or values["name"] or values["unit_barcode"])
            and (values["box_barcode"] or values["units_per_box"])
        )
        if is_continuation:
            prev = records[-1]
            records.append(
                Record(
                    row=i + 2,
                    group=prev.group,
                    ref=prev.ref,
                    name=prev.name,
                    unit_barcode=prev.unit_barcode,
                    box_barcode=values["box_barcode"],
                    units_per_box=values["units_per_box"],
                    uom=values["uom"] or prev.uom,
                    box_is_item=values["box_is_item"] or prev.box_is_item,
                    inherited=True,
                )
            )
        else:
            group += 1
            records.append(Record(row=i + 2, group=group, **values))

    used = {field: headers[idx] for field, idx in mapping.items()}
    return records, used
